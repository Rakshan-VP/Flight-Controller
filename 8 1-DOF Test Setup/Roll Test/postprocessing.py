# Values from both the sensors are matched by time stamping
# Done with almost same initial angle so that we can properly compare the results
# Test done for almost same time period
# PID gain values is almost constant so that the we can compare the performance of the filters properly. 

from openpyxl import load_workbook
import matplotlib.pyplot as plt
import os

def read_excel_columns(filename):
    wb = load_workbook(filename, data_only=True)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    data = {h: [] for h in headers}

    for row in ws.iter_rows(min_row=2, values_only=True):
        for h, v in zip(headers, row):
            data[h].append(v)

    return data


# ===== Get directory of this script =====
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

comp_file = os.path.join(
    BASE_DIR,
    "roll_pid_experiment_comp.xlsx"
)

madg_file = os.path.join(
    BASE_DIR,
    "roll_pid_experiment_madgwick.xlsx"
)

print("Loading:")
print(comp_file)
print(madg_file)

# ===== Load data =====
comp = read_excel_columns(comp_file)
madg = read_excel_columns(madg_file)

t_comp = range(len(comp["xsens_roll"]))
t_madg = range(len(madg["xsens_roll"]))

# ===== 2x2 PLOTS =====
fig, axs = plt.subplots(2, 2, figsize=(14, 10))

# ---- Plot 1: XSens vs Roll (Complementary) ----
axs[0, 0].plot(t_comp, comp["xsens_roll"], label="XSens", alpha=0.6)
axs[0, 0].plot(t_comp, comp["roll_comp"], label="Roll Comp", linewidth=2)
axs[0, 0].set_title("Complementary Filter Roll")
axs[0, 0].set_xlabel("Sample Index")
axs[0, 0].set_ylabel("Roll (deg)")
axs[0, 0].grid(True)
axs[0, 0].legend()

# ---- Plot 2: XSens vs Roll (Madgwick) ----
axs[0, 1].plot(t_madg, madg["xsens_roll"], label="XSens", alpha=0.6)
axs[0, 1].plot(t_madg, madg["roll_madgwick"], label="Roll Madgwick", linewidth=2)
axs[0, 1].set_title("Madgwick Filter Roll")
axs[0, 1].set_xlabel("Sample Index")
axs[0, 1].set_ylabel("Roll (deg)")
axs[0, 1].grid(True)
axs[0, 1].legend()

# ---- Plot 3: Motor PWMs (Complementary) ----
axs[1, 0].plot(t_comp, comp["left_front(25)"], label="LF 25")
axs[1, 0].plot(t_comp, comp["left_back(33)"], label="LB 33")
axs[1, 0].plot(t_comp, comp["right_front(32)"], label="RF 32")
axs[1, 0].plot(t_comp, comp["right_back(26)"], label="RB 26")
axs[1, 0].set_title("Motor Outputs (Comp)")
axs[1, 0].set_xlabel("Sample Index")
axs[1, 0].set_ylabel("PWM (µs)")
axs[1, 0].grid(True)
axs[1, 0].legend(ncol=2)

# ---- Plot 4: Motor PWMs (Madgwick) ----
axs[1, 1].plot(t_madg, madg["left_front(25)"], label="LF 25")
axs[1, 1].plot(t_madg, madg["left_back(33)"], label="LB 33")
axs[1, 1].plot(t_madg, madg["right_front(32)"], label="RF 32")
axs[1, 1].plot(t_madg, madg["right_back(26)"], label="RB 26")
axs[1, 1].set_title("Motor Outputs (Madgwick)")
axs[1, 1].set_xlabel("Sample Index")
axs[1, 1].set_ylabel("PWM (µs)")
axs[1, 1].grid(True)
axs[1, 1].legend(ncol=2)

plt.tight_layout()
plt.show()
